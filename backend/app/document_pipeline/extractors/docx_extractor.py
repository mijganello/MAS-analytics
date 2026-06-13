from __future__ import annotations
from docx import Document as DocxDocument
from app.document_pipeline.extractors.base import BaseExtractor, ExtractedDocument
from app.core.logging import logger


class DocxExtractor(BaseExtractor):
    @property
    def supported_types(self) -> list[str]:
        return ["docx"]

    def extract(self, file_path: str, file_id: str) -> ExtractedDocument:
        doc = DocxDocument(file_path)
        sections = []
        current_heading = "Introduction"

        for para in doc.paragraphs:
            if para.style.name.startswith("Heading"):
                current_heading = para.text
            if para.text.strip():
                sections.append({
                    "section_id": current_heading,
                    "content": para.text,
                    "style": para.style.name,
                    "type": "text",
                })

        raw_text = "\n".join(s["content"] for s in sections)
        logger.info("docx_extracted", file_id=file_id, paragraphs=len(sections))

        return ExtractedDocument(
            file_id=file_id,
            filename=file_path,
            file_type="docx",
            raw_text=raw_text,
            sections=sections,
            metadata={"paragraphs": len(sections)},
        )


class TxtExtractor(BaseExtractor):
    @property
    def supported_types(self) -> list[str]:
        return ["txt"]

    def extract(self, file_path: str, file_id: str) -> ExtractedDocument:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return ExtractedDocument(
            file_id=file_id,
            filename=file_path,
            file_type="txt",
            raw_text=content,
            sections=[{"section_id": "main", "content": content, "type": "text"}],
            metadata={"chars": len(content)},
        )
